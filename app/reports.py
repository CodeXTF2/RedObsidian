"""Report export: DOCX (optionally on an admin-uploaded Word template) and Markdown."""
import io
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from flask import current_app
from werkzeug.utils import secure_filename

from .maintenance import project_upload_dir
from .markdown_render import parse_blocks, parse_inline, safe_image
from .models import SEVERITIES, Upload

SEVERITY_COLORS = {"critical": "7B1FA2", "high": "D32F2F", "medium": "F57C00", "low": "1976D2", "info": "607D8B"}
DOCX_IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}
UPLOAD_URL_RE = re.compile(r"^(?:/p/\d+)?/uploads/([A-Za-z0-9._-]+)$")


def template_path():
    return os.path.join(current_app.instance_path, "report_template.docx")


def report_filename(project, extension):
    settings = project.report
    base = secure_filename(settings.get("title") or f"{project.name} report") or f"project-{project.id}-report"
    return f"{base}.{extension}"


def resolve_image(project, url):
    """Local path of an image upload in this project, or None (other projects' files are never embedded)."""
    match = UPLOAD_URL_RE.match(safe_image(url) or "")
    if not match:
        return None
    upload = Upload.query.filter_by(filename=match.group(1), project_id=project.id).first()
    if not upload or os.path.splitext(upload.filename)[1].lower() not in DOCX_IMAGE_TYPES:
        return None
    path = os.path.join(project_upload_dir(project.id, create=False), upload.filename)
    return path if os.path.isfile(path) else None


# --- DOCX ---------------------------------------------------------------------

def _style(document, name, fallback="Normal"):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles[fallback]


def _shade(cell, hex_color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    properties.append(shading)


def _add_runs(paragraph, text):
    for kind, value in parse_inline(text):
        if kind == "code":
            run = paragraph.add_run(value)
            run.font.name = "Consolas"
        elif kind == "bold":
            paragraph.add_run(value).bold = True
        elif kind == "italic":
            paragraph.add_run(value).italic = True
        elif kind == "link":
            label, url = value
            paragraph.add_run(f"{label} ({url})" if url.startswith("http") else label)
        elif kind == "image":
            paragraph.add_run(f"[{value[0] or 'image'}]")
        else:
            paragraph.add_run(value)


def _add_markdown(document, project, text, base_level=3):
    for kind, data in parse_blocks(text):
        if kind == "heading":
            document.add_heading(data["text"], level=min(9, base_level + data["level"] - 1))
        elif kind == "paragraph":
            _add_runs(document.add_paragraph(), data["text"])
        elif kind == "quote":
            _add_runs(document.add_paragraph(style=_style(document, "Quote")), data["text"])
        elif kind == "list":
            style = _style(document, "List Number" if data["ordered"] else "List Bullet")
            for item in data["items"]:
                _add_runs(document.add_paragraph(style=style), item)
        elif kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(data["text"])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            paragraph.paragraph_format.left_indent = Inches(0.2)
        elif kind == "image":
            path = resolve_image(project, data["url"])
            if path:
                try:
                    document.add_picture(path, width=Inches(6))
                except Exception:
                    document.add_paragraph(f"[image: {data['alt'] or os.path.basename(path)}]")
                    continue
                if data["alt"]:
                    caption = document.add_paragraph(data["alt"], style=_style(document, "Caption"))
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                document.add_paragraph(f"[image: {data['alt'] or data['url']}]")


def _section(document, project, heading, text):
    if text and text.strip():
        document.add_heading(heading, level=1)
        _add_markdown(document, project, text, base_level=2)


def _add_findings(document, project, context):
    document.add_heading("Findings summary", level=1)
    counts = context["counts"]
    table = document.add_table(rows=2, cols=len(SEVERITIES))
    table.style = _style(document, "Table Grid", "Normal Table")
    for column, level in enumerate(SEVERITIES):
        header = table.cell(0, column)
        header.text = level.capitalize()
        _shade(header, SEVERITY_COLORS[level])
        for run in header.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.bold = True
        table.cell(1, column).text = str(counts.get(level, 0))

    numbered = context["numbered"]
    if numbered:
        document.add_paragraph()
        overview = document.add_table(rows=1, cols=5)
        overview.style = _style(document, "Table Grid", "Normal Table")
        for column, label in enumerate(("ID", "Finding", "Severity", "CVSS", "Status")):
            overview.cell(0, column).text = label
            overview.cell(0, column).paragraphs[0].runs[0].bold = True
        for ref, finding in numbered:
            row = overview.add_row().cells
            row[0].text = ref
            row[1].text = finding.title
            row[2].text = finding.severity.capitalize()
            _shade(row[2], SEVERITY_COLORS.get(finding.severity, "FFFFFF"))
            row[3].text = f"{finding.cvss_score:.1f}" if finding.cvss_score is not None else "—"
            row[4].text = finding.status.replace("_", " ").capitalize()
    else:
        document.add_paragraph("No findings were recorded.")

    if numbered:
        document.add_page_break()
        document.add_heading("Detailed findings", level=1)
    for ref, finding in numbered:
        document.add_heading(f"{ref}: {finding.title}", level=2)
        meta = document.add_table(rows=0, cols=2)
        meta.style = _style(document, "Table Grid", "Normal Table")
        rows = [("Severity", finding.severity.capitalize()),
                ("CVSS 3.1", (f"{finding.cvss_score:.1f}" if finding.cvss_score is not None else "—")
                 + (f" ({finding.cvss_vector})" if finding.cvss_vector else "")),
                ("Status", finding.status.replace("_", " ").capitalize())]
        if finding.affected_assets:
            rows.append(("Affected assets", finding.affected_assets))
        for label, value in rows:
            cells = meta.add_row().cells
            cells[0].text = label
            cells[0].paragraphs[0].runs[0].bold = True
            cells[1].text = value
        _shade(meta.rows[0].cells[1], SEVERITY_COLORS.get(finding.severity, "FFFFFF"))
        for label, text in (("Description", finding.description), ("Impact", finding.impact),
                            ("Remediation", finding.remediation), ("References", finding.references)):
            if text:
                document.add_heading(label, level=3)
                _add_markdown(document, project, text, base_level=4)
        evidence = [context["pages"][node_id] for node_id in finding.evidence_node_ids if node_id in context["pages"]]
        if evidence:
            document.add_heading("Evidence", level=3)
            for page in evidence:
                document.add_heading(page.title, level=4)
                _add_markdown(document, project, page.notes or "", base_level=5)


def build_docx(project, context):
    settings = context["settings"]
    document = Document(template_path()) if os.path.isfile(template_path()) else Document()

    title = settings.get("title") or f"{project.name} — Security Assessment Report"
    document.add_paragraph(title, style=_style(document, "Title"))
    for label, key in (("Client", "client"), ("Assessment", "assessment_type"), ("Version", "version"),
                       ("Classification", "classification"), ("Prepared by", "prepared_by")):
        if settings.get(key):
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{label}: ").bold = True
            paragraph.add_run(settings[key])
    if settings.get("start_date") or settings.get("end_date"):
        paragraph = document.add_paragraph()
        paragraph.add_run("Testing window: ").bold = True
        paragraph.add_run(f"{settings.get('start_date') or '?'} to {settings.get('end_date') or '?'}")
    document.add_page_break()

    _section(document, project, "Executive summary", settings.get("executive_summary"))
    _section(document, project, "Scope", settings.get("scope"))
    _section(document, project, "Methodology", settings.get("methodology"))

    if context["findings_enabled"]:
        _add_findings(document, project, context)

    _section(document, project, "Conclusion", settings.get("conclusion"))

    if context["events"]:
        document.add_page_break()
        document.add_heading("Appendix: activity timeline", level=1)
        timeline = document.add_table(rows=1, cols=3)
        timeline.style = _style(document, "Table Grid", "Normal Table")
        for column, label in enumerate(("Time (UTC)", "Activity", "Details")):
            timeline.cell(0, column).text = label
            timeline.cell(0, column).paragraphs[0].runs[0].bold = True
        for event in context["events"]:
            cells = timeline.add_row().cells
            cells[0].text = event.occurred_at.strftime("%Y-%m-%d %H:%M")
            cells[1].text = event.title
            cells[2].text = event.body or ""

    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


# --- Markdown -------------------------------------------------------------------

def _markdown_findings(context):
    lines = ["## Findings summary", "", "| " + " | ".join(level.capitalize() for level in SEVERITIES) + " |",
              "|" + "---|" * len(SEVERITIES),
              "| " + " | ".join(str(context["counts"].get(level, 0)) for level in SEVERITIES) + " |", ""]
    if context["numbered"]:
        lines += ["| ID | Finding | Severity | CVSS | Status |", "|---|---|---|---|---|"]
        for ref, finding in context["numbered"]:
            score = f"{finding.cvss_score:.1f}" if finding.cvss_score is not None else "—"
            title = finding.title.replace("|", "\\|")
            lines.append(f"| {ref} | {title} | {finding.severity.capitalize()} | {score} | {finding.status} |")
        lines.append("")
        lines += ["## Detailed findings", ""]
    for ref, finding in context["numbered"]:
        lines += [f"### {ref}: {finding.title}", "", f"- **Severity:** {finding.severity.capitalize()}"]
        if finding.cvss_score is not None:
            lines.append(f"- **CVSS 3.1:** {finding.cvss_score:.1f}" + (f" (`{finding.cvss_vector}`)" if finding.cvss_vector else ""))
        lines.append(f"- **Status:** {finding.status}")
        if finding.affected_assets:
            lines.append(f"- **Affected assets:** {finding.affected_assets}")
        lines.append("")
        for label, text in (("Description", finding.description), ("Impact", finding.impact),
                            ("Remediation", finding.remediation), ("References", finding.references)):
            if text:
                lines += [f"#### {label}", "", text, ""]
        evidence = [context["pages"][node_id] for node_id in finding.evidence_node_ids if node_id in context["pages"]]
        if evidence:
            lines += ["#### Evidence", ""]
            for page in evidence:
                lines += [f"##### {page.title}", "", page.notes or "", ""]
    return lines


def build_markdown(project, context):
    settings = context["settings"]
    lines = [f"# {settings.get('title') or project.name + ' — Security Assessment Report'}", ""]
    for label, key in (("Client", "client"), ("Assessment", "assessment_type"), ("Version", "version"),
                       ("Classification", "classification"), ("Prepared by", "prepared_by")):
        if settings.get(key):
            lines.append(f"- **{label}:** {settings[key]}")
    if settings.get("start_date") or settings.get("end_date"):
        lines.append(f"- **Testing window:** {settings.get('start_date') or '?'} to {settings.get('end_date') or '?'}")
    lines.append("")
    for heading, key in (("Executive summary", "executive_summary"), ("Scope", "scope"), ("Methodology", "methodology")):
        if settings.get(key):
            lines += [f"## {heading}", "", settings[key], ""]

    if context["findings_enabled"]:
        lines += _markdown_findings(context)
    if settings.get("conclusion"):
        lines += ["## Conclusion", "", settings["conclusion"], ""]
    if context["events"]:
        lines += ["## Appendix: activity timeline", "", "| Time (UTC) | Activity |", "|---|---|"]
        for event in context["events"]:
            title = event.title.replace("|", "\\|")
            lines.append(f"| {event.occurred_at:%Y-%m-%d %H:%M} | {title} |")
    return "\n".join(lines).rstrip() + "\n"
