import io
import json
import zipfile

import pytest
from docx import Document

from app import db
from app.findings import cvss31_score, severity_for_score
from app.markdown_render import to_html
from app.models import Finding, GraphNode


@pytest.mark.parametrize("vector, score", [
    ("CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H", 9.8),
    ("CVSS:3.1/AV:N/AC:L/PR:N/UI:R/S:C/C:L/I:L/A:N", 6.1),
    ("CVSS:3.1/AV:L/AC:H/PR:H/UI:R/S:U/C:L/I:N/A:N", 1.8),
    ("CVSS:3.1/AV:N/AC:L/PR:L/UI:N/S:C/C:H/I:H/A:H", 9.9),
    ("CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:N/I:N/A:N", 0.0),
])
def test_cvss_scores(vector, score):
    assert cvss31_score(vector) == score


def test_cvss_rejects_incomplete_vectors():
    assert cvss31_score("CVSS:3.1/AV:N/AC:L") is None
    assert cvss31_score("nonsense") is None
    assert severity_for_score(9.8) == "critical" and severity_for_score(5) == "medium" and severity_for_score(0) == "info"


def test_markdown_escapes_payloads():
    html = to_html('<img src=x onerror=alert(1)>\n\n[click](javascript:alert(1))\n\n![a](https://evil/x.png)\n\n```\n<script>\n```')
    assert "<img src=x" not in html and "&lt;img" in html
    assert "javascript:" not in html
    assert "https://evil" not in html  # only same-site images
    assert "&lt;script&gt;" in html


def create_finding(client, project_id, **fields):
    data = {"title": "SQL injection", "status": "open", "severity": "", "cvss_vector": "",
            "description": "Found it", **fields}
    return client.post(f"/p/{project_id}/findings", data=data)


def test_finding_crud_and_cvss(app, make_user, make_project, login_as):
    alice = make_user("alice")
    project_id = make_project("A", {alice: "owner"}, findings=True)
    client = login_as("alice")
    response = create_finding(client, project_id, cvss_vector="CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H")
    assert response.status_code == 302
    with app.app_context():
        finding = Finding.query.one()
        assert finding.cvss_score == 9.8 and finding.severity == "critical"
        finding_id = finding.id
    page = client.get(f"/p/{project_id}/findings").get_data(as_text=True)
    assert "SQL injection" in page
    client.post(f"/p/{project_id}/findings/{finding_id}", data={"title": "SQLi (fixed)", "status": "remediated",
                                                                "severity": "high", "cvss_vector": ""})
    with app.app_context():
        finding = db.session.get(Finding, finding_id)
        assert (finding.title, finding.status, finding.severity) == ("SQLi (fixed)", "remediated", "high")
    client.post(f"/p/{project_id}/findings/{finding_id}/delete")
    with app.app_context():
        assert Finding.query.count() == 0


def test_findings_are_project_scoped(app, make_user, make_project, login_as):
    alice, bob = make_user("alice"), make_user("bob")
    project_a = make_project("A", {alice: "owner"}, findings=True)
    project_b = make_project("B", {bob: "owner"}, findings=True)
    create_finding(login_as("alice"), project_a)
    with app.app_context():
        finding_id = Finding.query.one().id
    bob_client = login_as("bob")
    assert bob_client.get(f"/p/{project_b}/findings/{finding_id}").status_code == 404
    assert bob_client.post(f"/p/{project_b}/findings/{finding_id}/delete").status_code == 404


def test_report_exports(app, make_user, make_project, login_as):
    alice = make_user("alice")
    project_id = make_project("A", {alice: "owner"}, findings=True, report=True)
    client = login_as("alice")
    with app.app_context():
        page = GraphNode.query.filter_by(project_id=project_id).first()
        page.title, page.notes = "Evidence page", "# Proof\n\n```\nid; whoami\n```"
        db.session.commit()
        page_id = page.id
    create_finding(client, project_id, cvss_vector="CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H",
                   description="Payload: <script>alert(1)</script>", evidence=str(page_id))
    create_finding(client, project_id, title="Draft thing", status="draft")
    client.post(f"/p/{project_id}/report", data={"title": "ACME Pentest", "client": "ACME",
                                                 "executive_summary": "Things were **bad**."})

    docx_response = client.get(f"/p/{project_id}/report.docx")
    assert docx_response.status_code == 200
    document = Document(io.BytesIO(docx_response.data))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "ACME Pentest" in text and "F-01: SQL injection" in text and "Evidence page" in text
    assert "Draft thing" not in text  # drafts are excluded by default

    markdown = client.get(f"/p/{project_id}/report.md").get_data(as_text=True)
    assert "### F-01: SQL injection" in markdown and "id; whoami" in markdown

    preview = client.get(f"/p/{project_id}/report/preview").get_data(as_text=True)
    assert "<script>alert(1)</script>" not in preview and "&lt;script&gt;" in preview


def test_report_template_is_used(app, make_user, make_project, login_as):
    make_user("root", is_admin=True)
    root_id = make_user("alice")
    project_id = make_project("A", {root_id: "owner"}, findings=True, report=True)
    template = Document()
    template.add_paragraph("ORGANISATION COVER")
    buffer = io.BytesIO()
    template.save(buffer)
    admin = login_as("root")
    response = admin.post("/admin/report-template", data={"template": (io.BytesIO(buffer.getvalue()), "t.docx")},
                          content_type="multipart/form-data")
    assert response.status_code == 302
    document = Document(io.BytesIO(login_as("alice").get(f"/p/{project_id}/report.docx").data))
    assert document.paragraphs[0].text == "ORGANISATION COVER"
    # A non-docx upload is refused
    bad = admin.post("/admin/report-template", data={"template": (io.BytesIO(b"not a zip"), "t.docx")},
                     content_type="multipart/form-data")
    assert bad.status_code == 302
    assert zipfile.is_zipfile(io.BytesIO(login_as("alice").get(f"/p/{project_id}/report.docx").data))


def test_report_settings_saved(app, make_user, make_project, login_as):
    alice = make_user("alice")
    project_id = make_project("A", {alice: "owner"}, findings=True, report=True)
    client = login_as("alice")
    client.post(f"/p/{project_id}/report", data={"client": "ACME", "include_timeline": "on"})
    from app.models import Project
    with app.app_context():
        report = json.loads(db.session.get(Project, project_id).report_json)
    assert report["client"] == "ACME" and report["include_timeline"] is True and report["include_drafts"] is False


def test_findings_off_by_default_and_owner_toggle(app, make_user, make_project, login_as):
    alice, ed = make_user("alice"), make_user("ed")
    project_id = make_project("A", {alice: "owner", ed: "editor"}, report=True)
    owner = login_as("alice")
    assert owner.get(f"/p/{project_id}/findings").status_code == 404
    assert b"Findings <span class=\"beta-tag\">Beta</span>" not in owner.get(f"/p/{project_id}/notes").data
    assert b"Findings summary" not in owner.get(f"/p/{project_id}/report/preview").data
    # Only owners can turn it on
    assert login_as("ed").post(f"/p/{project_id}/features", data={"feature": "findings", "findings_enabled": "on"}).status_code == 403
    owner.post(f"/p/{project_id}/features", data={"feature": "findings", "findings_enabled": "on"})
    assert owner.get(f"/p/{project_id}/findings").status_code == 200
    assert b"beta-tag" in owner.get(f"/p/{project_id}/notes").data
    create_finding(owner, project_id)
    owner.post(f"/p/{project_id}/features", data={"feature": "findings"})
    assert owner.get(f"/p/{project_id}/findings").status_code == 404
    # Turning it off hides findings from reports but keeps them
    assert "SQL injection" not in owner.get(f"/p/{project_id}/report.md").get_data(as_text=True)
    with app.app_context():
        assert Finding.query.count() == 1


def test_report_off_by_default_and_owner_toggle(app, make_user, make_project, login_as):
    alice, ed = make_user("alice"), make_user("ed")
    project_id = make_project("A", {alice: "owner", ed: "editor"}, findings=True)
    owner = login_as("alice")
    for path in ("report", "report/preview", "report.docx", "report.md"):
        assert owner.get(f"/p/{project_id}/{path}").status_code == 404, path
    assert owner.post(f"/p/{project_id}/report", data={"client": "ACME"}).status_code == 404
    assert b"Report <span class=\"beta-tag\">Beta</span>" not in owner.get(f"/p/{project_id}/notes").data
    assert login_as("ed").post(f"/p/{project_id}/features", data={"feature": "report", "report_enabled": "on"}).status_code == 403
    assert owner.post(f"/p/{project_id}/features", data={"feature": "nope"}).status_code == 400
    owner.post(f"/p/{project_id}/features", data={"feature": "report", "report_enabled": "on"})
    assert b"Report <span class=\"beta-tag\">Beta</span>" in owner.get(f"/p/{project_id}/notes").data
    owner.post(f"/p/{project_id}/report", data={"client": "ACME"})
    assert owner.get(f"/p/{project_id}/report.md").status_code == 200
    # Toggling Report leaves Findings alone, and turning it off keeps the report details
    owner.post(f"/p/{project_id}/features", data={"feature": "report"})
    assert owner.get(f"/p/{project_id}/report").status_code == 404
    assert owner.get(f"/p/{project_id}/findings").status_code == 200
    from app.models import Project
    with app.app_context():
        assert json.loads(db.session.get(Project, project_id).report_json)["client"] == "ACME"
